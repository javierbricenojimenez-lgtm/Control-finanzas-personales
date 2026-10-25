
import streamlit as st
import pandas as pd
from config import MODO_PUBLICO
from demo_sesion import crear_base_demo_temporal, eliminar_base_demo_temporal
from pathlib import Path
from io import BytesIO
from repositorio import (
    obtener_movimientos,
    agregar_movimiento,
    actualizar_movimiento,
    eliminar_movimiento_sql
)

# Formato de pesos chilenos
def formato_clp(valor):
    monto = f"{abs(valor):,.0f}".replace(",", ".")

    if valor < 0:
        return f"-${monto}"

    return f"${monto}"

# Configuracion de la pagina
st.set_page_config(
    page_title="Control Financiero Personal",
    page_icon=":material/account_balance_wallet:",
    layout="wide"
)

st.sidebar.title(":material/account_balance_wallet: Control Financiero Personal")
st.sidebar.caption("FINANZAS PERSONALES · 2.0")
st.title("Tu dinero, con claridad.")
st.caption("Organiza tus movimientos y entiende en qué estás gastando.")
if "mensaje_edicion" in st.session_state:
    st.success(st.session_state.pop("mensaje_edicion"))
if "mensaje_eliminacion" in st.session_state:
    st.success(st.session_state.pop("mensaje_eliminacion"))
# 1. Obtener datos desde SQLite

# Elegir origen de datos
if MODO_PUBLICO:
    # Cada sesión obtiene una base temporal independiente
    if "base_demo_sesion" not in st.session_state:
        st.session_state["base_demo_sesion"] = crear_base_demo_temporal()

    base_datos_actual = st.session_state["base_demo_sesion"]
    modo_demo = True

    st.sidebar.badge("Demostración pública", icon=":material/science:", color="green")
    st.sidebar.caption("Los datos son ficticios.")

    if st.sidebar.button("Reiniciar demostración"):
        ruta_anterior = st.session_state["base_demo_sesion"]
        ruta_nueva = crear_base_demo_temporal()

        st.session_state["base_demo_sesion"] = ruta_nueva
        eliminar_base_demo_temporal(ruta_anterior)

        st.rerun()

else:
    modo_demo = st.sidebar.toggle(
        "🧪 Modo demostración",
        value=False
    )

    if modo_demo:
        if "base_demo_sesion" not in st.session_state:
            st.session_state["base_demo_sesion"] = crear_base_demo_temporal()
        base_datos_actual = st.session_state["base_demo_sesion"]
        st.sidebar.info("Estás utilizando datos ficticios.")
    else:
        # Restrict personal access through the actual server bind address.
        # This is local isolation, not a replacement for remote authentication.
        if st.get_option("server.address") not in {"127.0.0.1", "::1"}:
            st.error("Para abrir tus datos personales, inicia la aplicación con iniciar_personal.ps1. El acceso personal está limitado a este computador.")
            st.stop()
        base_datos_actual = str(Path(__file__).parent / "finanzas.db")
        if not Path(base_datos_actual).is_file():
            st.error("Esta copia no contiene tu base personal. Utiliza el modo demostración para probarla.")
            st.stop()
        st.sidebar.info("Estás utilizando tu base personal.")

movimientos = obtener_movimientos(base_datos_actual)
etiquetas_movimientos = {
    m["id"]: f"#{m['id']} · {m['tipo']} · {formato_clp(m['cantidad'])} · {m['categoria']}"
    for m in movimientos
}

if not movimientos:
    st.info("Todavía no hay movimientos registrados.")

df = pd.DataFrame(movimientos, columns=["id", "tipo", "cantidad", "categoria", "fecha"])

# 2. Preparar fechas y periodos
df["periodo"] = df["fecha"].apply(
    lambda fecha: str(fecha)[:7]
    if pd.notna(fecha) and fecha != "Sin fecha"
    else "Sin fecha"
)

seccion = st.sidebar.radio(
    "Menú", ["Resumen", "Registrar movimiento", "Movimientos", "Categorías", "Gráficos", "Reportes"],
    key="seccion",
    label_visibility="collapsed"
)
st.header(seccion)
descripciones = {
    "Resumen": "Una mirada a tus ingresos, gastos y saldo del período.",
    "Registrar movimiento": "Anota un ingreso o un gasto en unos pocos pasos.",
    "Movimientos": "Revisa tu historial y mantén tus registros al día.",
    "Categorías": "Descubre cómo se distribuyen tus gastos.",
    "Gráficos": "Encuentra patrones en tus finanzas.",
    "Reportes": "Lleva tus movimientos a Excel o descarga un archivo CSV."
}
st.caption(descripciones[seccion])

# 3. Filtros del dashboard
categoria_seleccionada = "Todas"
periodo_seleccionado = "Todos"
if seccion != "Registrar movimiento":
    st.caption(":material/tune: Ajusta la vista")

    categorias = sorted(df["categoria"].unique().tolist())
    periodos = sorted(df["periodo"].unique().tolist())

    col_filtro1, col_filtro2 = st.columns(2)

    with col_filtro1:
        categoria_seleccionada = st.selectbox(
            "Filtrar por categoría",
            ["Todas"] + categorias
        )

    with col_filtro2:
        periodo_seleccionado = st.selectbox(
            "Filtrar por mes",
            ["Todos"] + periodos
        )

# 4. Aplicar filtros
movimientos_filtrados = df.copy()

if categoria_seleccionada != "Todas":
    movimientos_filtrados = movimientos_filtrados[
        movimientos_filtrados["categoria"]
        == categoria_seleccionada
    ]

if periodo_seleccionado != "Todos":
    movimientos_filtrados = movimientos_filtrados[
        movimientos_filtrados["periodo"]
        == periodo_seleccionado
    ]

# 5. Calcular indicadores filtrados
ingresos = movimientos_filtrados.loc[
    movimientos_filtrados["tipo"] == "Ingreso",
    "cantidad"
].sum()

gastos = movimientos_filtrados.loc[
    movimientos_filtrados["tipo"] == "Gasto",
    "cantidad"
].sum()

saldo = ingresos - gastos

gastos_grafico = movimientos_filtrados[
    movimientos_filtrados["tipo"] == "Gasto"
]

resumen_grafico = (
    gastos_grafico.groupby("categoria")["cantidad"]
    .sum()
    .sort_values(ascending=False)
)

if seccion == 'Resumen':
    st.subheader("Tu balance")

    col1, col2, col3 = st.columns(3)

    col1.metric("Ingresos", formato_clp(ingresos), border=True)
    col2.metric("Gastos", formato_clp(gastos), border=True)
    col3.metric("Saldo del período", formato_clp(saldo), border=True)

    st.caption("Los indicadores reflejan los filtros seleccionados; el saldo del período no es el saldo de una cuenta bancaria.")
    izquierda, derecha = st.columns([1.6, 1])
    with izquierda:
        with st.container(border=True):
            st.subheader("¿Dónde se va tu dinero?")
            if resumen_grafico.empty:
                st.info("No hay gastos para los filtros seleccionados.")
            else:
                st.bar_chart(resumen_grafico, color="#087F70", height=290, x_label="Categoría", y_label="Monto (CLP)", alt="Gastos por categoría para los filtros seleccionados")
    with derecha:
        with st.container(border=True):
            st.subheader("En este período")
            if gastos > ingresos:
                st.warning("Los gastos superan los ingresos de esta selección.")
            st.metric("Movimientos", len(movimientos_filtrados))
            if ingresos > 0:
                proporcion = min(float(gastos / ingresos), 1.0)
                st.progress(proporcion, text=f"Gastos: {gastos / ingresos:.0%} de los ingresos")
            if not resumen_grafico.empty:
                st.caption("Categoría con mayor gasto")
                st.write(f"**{resumen_grafico.index[0]}** · {formato_clp(resumen_grafico.iloc[0])}")
    with st.container(border=True):
        st.subheader("Últimos movimientos")
        recientes = movimientos_filtrados.sort_values(["fecha", "id"], ascending=False).head(5)
        st.dataframe(recientes[["fecha", "tipo", "categoria", "cantidad"]], hide_index=True,
                     column_config={"fecha": "Fecha", "tipo": "Tipo", "categoria": "Categoría", "cantidad": st.column_config.NumberColumn("Monto (CLP)", format="$ %.0f")}, width="stretch")

if seccion == 'Registrar movimiento':
    formulario, ayuda = st.columns([1.7, 1], vertical_alignment="top")
    with ayuda:
        with st.container(border=True):
            st.subheader(":material/tips_and_updates: Antes de guardar")
            st.write("**Ingreso:** dinero que recibes, como sueldo, ventas o bonos.")
            st.write("**Gasto:** dinero que utilizas en compras, servicios u otras categorías.")
            st.caption("Ingresa el monto en pesos chilenos y selecciona la fecha real del movimiento.")
        if modo_demo:
            st.info("Estás practicando con datos ficticios. Puedes reiniciar la demostración cuando quieras.", icon=":material/science:")
        else:
            st.info("Los movimientos se guardan en tu base personal.", icon=":material/database:")
    with formulario:
        st.subheader("Detalles del movimiento")

        # Mostrar confirmacion despues de guardar
        if "registro_exitoso" in st.session_state:
            st.success(st.session_state.pop("registro_exitoso"))

        # Elegir tipo antes de abrir el formulario
        tipo_nuevo = st.radio(
            "Tipo de movimiento",
            ["Gasto", "Ingreso"],
            horizontal=True
        )

        with st.form("formulario_movimiento", clear_on_submit=True):

            cantidad_nueva = st.number_input(
                "Monto ($ CLP)",
                min_value=1,
                step=1000
            )

            # Categorias de gastos
            categorias_gastos = [
                "Alimentación",
                "Transporte",
                "Salud",
                "Vivienda",
                "Servicios",
                "Entretenimiento",
                "Educación",
                "Otros"
            ]

            # Categorias de ingresos
            categorias_ingresos = [
                "Sueldo",
                "Ventas",
                "Bonos",
                "Inversiones",
                "Trabajos independientes",
                "Otros ingresos"
            ]

            # Seleccionar categorias segun el tipo
            if tipo_nuevo == "Gasto":
                categorias_disponibles = categorias_gastos
            else:
                categorias_disponibles = categorias_ingresos

            categoria_nueva = st.selectbox(
                "Categoría",
                categorias_disponibles
            )

            fecha_nueva = st.date_input(
                "Fecha del movimiento"
            )

            guardar = st.form_submit_button("Guardar movimiento", type="primary", width="stretch")

        # Guardar el movimiento en SQLite
        if guardar:
            if not categoria_nueva.strip():
                st.warning("Debes ingresar una categoría.")
            else:
    
                nuevo_id = agregar_movimiento(
                     tipo_nuevo,
                     cantidad_nueva,
                     categoria_nueva.strip(),
                     fecha_nueva.isoformat(),
                     base_datos=base_datos_actual
        )

                st.session_state["registro_exitoso"] = (
                    f"✅ ¡Registro completado correctamente! "
                    f"Movimiento #{nuevo_id} guardado."
                )

                st.rerun()

if seccion == 'Movimientos':
    st.divider()
    st.subheader("Historial")

    st.dataframe(
        movimientos_filtrados,
        column_config={"id": "ID", "tipo": "Tipo", "cantidad": st.column_config.NumberColumn("Monto (CLP)", format="$ %.0f"), "categoria": "Categoría", "fecha": "Fecha", "periodo": "Período"},
        width="stretch",
        hide_index=True
    )

    st.caption(
        f"Movimientos encontrados: {len(movimientos_filtrados)}"
    )

    with st.expander("Editar un movimiento", icon=":material/edit:"):
        st.subheader("Editar un movimiento")

        if not movimientos_filtrados.empty:

            ids_disponibles = movimientos_filtrados["id"].tolist()

            id_seleccionado = st.selectbox(
                "Selecciona el ID del movimiento",
                ids_disponibles,
                format_func=lambda identificador: etiquetas_movimientos[identificador]
            )

            movimiento = movimientos_filtrados[
                movimientos_filtrados["id"] == id_seleccionado
            ].iloc[0]

            st.write("Tipo:", movimiento["tipo"])
            st.write("Monto actual:", formato_clp(movimiento["cantidad"]))
            st.write("Categoría actual:", movimiento["categoria"])
            st.write("Fecha:", movimiento["fecha"])

            with st.form("formulario_edicion"):

                monto_editado = st.number_input(
                    "Nuevo monto ($ CLP)",
                    min_value=1,
                    value=int(movimiento["cantidad"]),
                    key=f"monto_editar_{id_seleccionado}",
                    step=1000
                )

                categoria_editada = st.text_input(
                    "Nueva categoría",
                    value=movimiento["categoria"],
                    key=f"categoria_editar_{id_seleccionado}"
                )

                guardar_cambios = st.form_submit_button(
                    "Guardar cambios"
                )

            if guardar_cambios:

                if not categoria_editada.strip():
                    st.warning("Debes ingresar una categoría.")
                else:
            
                    filas = actualizar_movimiento(
                         int(id_seleccionado),
                         monto_editado,
                         categoria_editada.strip(),
                         base_datos=base_datos_actual
        )

                    if filas > 0:
                        st.session_state["mensaje_edicion"] = (
                            "✅ Movimiento actualizado correctamente."
                        )
                        st.rerun()
                    else:
                        st.error("No se pudo actualizar el movimiento.")

        else:
            st.info("No hay movimientos disponibles para editar.")

    # 10. Eliminar movimientos
    with st.expander("Eliminar un movimiento", icon=":material/delete:"):
        st.subheader("Eliminar un movimiento")

        if not movimientos_filtrados.empty:

            id_eliminar = st.selectbox(
                "Selecciona el ID que deseas eliminar",
                movimientos_filtrados["id"].tolist(),
                key="id_eliminar",
                format_func=lambda identificador: etiquetas_movimientos[identificador]
            )

            seleccionado = movimientos_filtrados.loc[movimientos_filtrados["id"] == id_eliminar].iloc[0]
            st.write(f"**{seleccionado['tipo']} · {formato_clp(seleccionado['cantidad'])}**")
            st.caption(f"{seleccionado['categoria']} · {seleccionado['fecha']}")
            st.caption("Esta acción elimina el registro seleccionado de la base actual.")
            confirmar_eliminacion = st.checkbox(
                f"Confirmo que quiero eliminar el movimiento #{id_eliminar}",
                key=f"confirmar_{id_eliminar}"
            )

            if st.button(
                "Eliminar definitivamente",
                disabled=not confirmar_eliminacion
            ):
                filas = eliminar_movimiento_sql(
                     int(id_eliminar),
                      base_datos=base_datos_actual
                    )

                if filas > 0:
                    st.session_state["mensaje_eliminacion"] = (
                        f"Movimiento #{id_eliminar} eliminado correctamente."
                    )
                    st.rerun()
                else:
                    st.error("No se pudo eliminar el movimiento.")

        else:
            st.info("No hay movimientos disponibles para eliminar.")

if seccion == 'Categorías':
    st.subheader("Resumen de gastos por categoría")
    if resumen_grafico.empty:
        st.info("No hay gastos para los filtros seleccionados.")
    else:
        tabla = resumen_grafico.rename("Total ($ CLP)").to_frame()
        tabla["Porcentaje (%)"] = (tabla["Total ($ CLP)"] / tabla["Total ($ CLP)"].sum() * 100).round(2)
        primero, segundo = st.columns(2)
        primero.metric("Total de gastos", formato_clp(tabla["Total ($ CLP)"].sum()), border=True)
        segundo.metric("Categorías con gastos", len(tabla), border=True)
        st.dataframe(tabla, width="stretch", column_config={
            "Total ($ CLP)": st.column_config.NumberColumn("Total ($ CLP)", format="$ %.0f"),
            "Porcentaje (%)": st.column_config.ProgressColumn("Participación (%)", min_value=0, max_value=100, format="%.2f%%")
        })

if seccion == 'Gráficos':
    st.divider()
    st.subheader("Gastos por categoría")

    if not resumen_grafico.empty:
        with st.container(border=True):
            st.bar_chart(resumen_grafico, color="#087F70", height=360, x_label="Categoría", y_label="Monto (CLP)", alt="Distribución de gastos por categoría")
        st.caption(f"Total de gastos de esta selección: {formato_clp(resumen_grafico.sum())}. Los ingresos no forman parte de este gráfico.")
    else:
        st.info("No hay gastos para los filtros seleccionados.")

if seccion == 'Reportes':
    st.divider()
    st.subheader("Tus datos, listos para descargar")
    st.caption(f"Período: {periodo_seleccionado} · Categoría: {categoria_seleccionada} · Moneda: CLP")

    st.info(f"El reporte incluirá {len(movimientos_filtrados)} movimientos, según los filtros seleccionados.", icon=":material/download:")

    # Preparar datos para exportar
    reporte = movimientos_filtrados.copy()

    # Archivo CSV
    csv = reporte.to_csv(
        index=False,
        sep=";"
    ).encode("utf-8-sig")

    # Archivo Excel
    archivo_excel = BytesIO()

    with pd.ExcelWriter(
        archivo_excel,
        engine="openpyxl"
    ) as writer:

        reporte.to_excel(
            writer,
            index=False,
            sheet_name="Movimientos"
        )

        hoja = writer.sheets["Movimientos"]

        # Ajustar ancho de columnas
        hoja.column_dimensions["A"].width = 10
        hoja.column_dimensions["B"].width = 15
        hoja.column_dimensions["C"].width = 18
        hoja.column_dimensions["D"].width = 25
        hoja.column_dimensions["E"].width = 18
        hoja.column_dimensions["F"].width = 15

        # Destacar encabezados
        from openpyxl.styles import Font, PatternFill

        for celda in hoja[1]:
            celda.font = Font(
                bold=True,
                color="FFFFFF"
            )
            celda.fill = PatternFill(
                start_color="1F4E78",
                end_color="1F4E78",
                fill_type="solid"
            )

        # Formato de montos en pesos chilenos
        for fila in range(2, hoja.max_row + 1):
            hoja[f"C{fila}"].number_format = '"$"#,##0'

        # Mantener encabezados visibles
        hoja.freeze_panes = "A2"

        # Activar filtros en Excel
        hoja.auto_filter.ref = hoja.dimensions

    # Botones de descarga
    col_csv, col_excel = st.columns(2)

    with col_csv:
        st.subheader(":material/table_view: Archivo CSV")
        st.caption("Formato de datos con separador punto y coma. Compatible con Excel.")
        st.download_button(
            label="Descargar CSV",
            data=csv,
            file_name="reporte_movimientos.csv",
            mime="text/csv",
            width="stretch"
        )

    with col_excel:
        st.subheader(":material/description: Libro de Excel")
        st.caption("Tabla con encabezados, filtros y formato monetario listo para revisar.")
        st.download_button(
            label="Descargar Excel",
            data=archivo_excel.getvalue(),
            file_name="reporte_movimientos.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch"
        )
